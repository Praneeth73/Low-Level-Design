"""Turn uploaded resumes and job-posting URLs into plain text."""
import io
import ipaddress
import socket
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

MAX_JD_CHARS = 30_000
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"
)


class ExtractionError(Exception):
    pass


def resume_to_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ExtractionError("Unsupported resume format. Upload a PDF, DOCX, TXT or MD file.")

    text = text.strip()
    if len(text) < 50:
        raise ExtractionError(
            "Could not read text from the resume. If it is a scanned PDF, upload a DOCX or text version."
        )
    return text


def _assert_public_host(url: str) -> None:
    """Refuse URLs that point at localhost or private networks (SSRF guard)."""
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise ExtractionError("Job link must be a full http(s) URL.")
    try:
        infos = socket.getaddrinfo(parsed.hostname, None)
    except socket.gaierror:
        raise ExtractionError(f"Could not resolve host {parsed.hostname}.")
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise ExtractionError("Job link points to a private network address.")


def job_description_from_url(url: str) -> str:
    _assert_public_host(url)
    try:
        resp = httpx.get(
            url,
            headers={"User-Agent": USER_AGENT, "Accept-Language": "en-US,en;q=0.9"},
            follow_redirects=True,
            timeout=20,
        )
    except httpx.HTTPError as e:
        raise ExtractionError(f"Could not fetch the job link: {e}")
    if resp.status_code >= 400:
        raise ExtractionError(
            f"The job site returned HTTP {resp.status_code}. "
            "Some sites (e.g. LinkedIn) block automated access - paste the description instead."
        )
    _assert_public_host(str(resp.url))

    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "header", "footer", "nav", "form"]):
        tag.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    lines = [line.strip() for line in main.get_text("\n").splitlines()]
    text = "\n".join(line for line in lines if line)
    if len(text) < 200:
        raise ExtractionError(
            "The job page had too little readable text (it may need JavaScript or a login). "
            "Paste the job description instead."
        )
    return text[:MAX_JD_CHARS]
