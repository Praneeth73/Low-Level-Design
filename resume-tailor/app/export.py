"""Render a Resume to DOCX or PDF."""
import io
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from .models import Resume

ACCENT = "1F3A5F"


def _contact_line(r: Resume) -> str:
    c = r.contact
    return "  |  ".join(v for v in [c.email, c.phone, c.location, c.linkedin, c.github, c.website] if v)


def _dates(start, end) -> str:
    return " - ".join(d for d in [start, end] if d)


# ---------------------------------------------------------------- DOCX

def _add_rule(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": ACCENT}.items():
        bottom.set(qn(k), v)
    border.append(bottom)
    p_pr.append(border)


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    _add_rule(p)


def _title_line(doc, left: str, right: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(doc.sections[0].page_width
                                              - doc.sections[0].left_margin
                                              - doc.sections[0].right_margin,
                                              alignment=2)  # right-aligned tab
    p.add_run(left).bold = True
    if right:
        r = p.add_run("\t" + right)
        r.italic = True


def _bullets(doc, items) -> None:
    for b in items:
        p = doc.add_paragraph(b, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)


def to_docx(r: Resume) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(42)
        s.left_margin = s.right_margin = Pt(50)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(r.name)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    if r.headline:
        p = doc.add_paragraph(r.headline)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact = _contact_line(r)
    if contact:
        p = doc.add_paragraph(contact)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size = Pt(9.5)

    if r.summary:
        _heading(doc, "Summary")
        doc.add_paragraph(r.summary)

    if r.skills:
        _heading(doc, "Skills")
        for g in r.skills:
            p = doc.add_paragraph()
            p.add_run(f"{g.category}: ").bold = True
            p.add_run(", ".join(g.skills))

    if r.experience:
        _heading(doc, "Experience")
        for e in r.experience:
            _title_line(doc, f"{e.title} - {e.company}", _dates(e.start_date, e.end_date))
            if e.location:
                doc.add_paragraph(e.location).runs[0].italic = True
            _bullets(doc, e.bullets)

    if r.projects:
        _heading(doc, "Projects")
        for pr in r.projects:
            tech = f" ({', '.join(pr.technologies)})" if pr.technologies else ""
            _title_line(doc, pr.name + tech, "")
            if pr.description:
                doc.add_paragraph(pr.description)
            _bullets(doc, pr.bullets)

    if r.education:
        _heading(doc, "Education")
        for ed in r.education:
            _title_line(doc, f"{ed.degree} - {ed.institution}", _dates(ed.start_date, ed.end_date))
            if ed.details:
                doc.add_paragraph(ed.details)

    if r.certifications:
        _heading(doc, "Certifications")
        _bullets(doc, r.certifications)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

def to_pdf(r: Resume) -> bytes:
    accent = HexColor("#" + ACCENT)
    name_style = ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=20, leading=24,
                                alignment=TA_CENTER, textColor=accent)
    center = ParagraphStyle("center", fontName="Helvetica", fontSize=9.5, leading=13, alignment=TA_CENTER)
    h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=accent,
                       spaceBefore=8, spaceAfter=1)
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=13)
    role = ParagraphStyle("role", parent=body, fontName="Helvetica-Bold", spaceBefore=4)
    sub = ParagraphStyle("sub", parent=body, fontName="Helvetica-Oblique", textColor=HexColor("#555555"))

    e = escape
    story = [Paragraph(e(r.name), name_style)]
    if r.headline:
        story.append(Paragraph(e(r.headline), center))
    contact = _contact_line(r)
    if contact:
        story.append(Paragraph(e(contact), center))

    def section(title):
        story.append(Paragraph(title.upper(), h))
        story.append(HRFlowable(width="100%", thickness=0.8, color=accent, spaceAfter=4))

    def bullets(items):
        if items:
            story.append(ListFlowable(
                [ListItem(Paragraph(e(b), body), leftIndent=12) for b in items],
                bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8,
            ))

    def title_line(left, right):
        text = f"{e(left)}"
        if right:
            text += f"<font name='Helvetica-Oblique' color='#555555'>  ({e(right)})</font>"
        story.append(Paragraph(text, role))

    if r.summary:
        section("Summary")
        story.append(Paragraph(e(r.summary), body))
    if r.skills:
        section("Skills")
        for g in r.skills:
            story.append(Paragraph(f"<b>{e(g.category)}:</b> {e(', '.join(g.skills))}", body))
    if r.experience:
        section("Experience")
        for x in r.experience:
            title_line(f"{x.title} - {x.company}", _dates(x.start_date, x.end_date))
            if x.location:
                story.append(Paragraph(e(x.location), sub))
            bullets(x.bullets)
    if r.projects:
        section("Projects")
        for pr in r.projects:
            tech = f" ({', '.join(pr.technologies)})" if pr.technologies else ""
            title_line(pr.name + tech, "")
            if pr.description:
                story.append(Paragraph(e(pr.description), body))
            bullets(pr.bullets)
    if r.education:
        section("Education")
        for ed in r.education:
            title_line(f"{ed.degree} - {ed.institution}", _dates(ed.start_date, ed.end_date))
            if ed.details:
                story.append(Paragraph(e(ed.details), body))
    if r.certifications:
        section("Certifications")
        bullets(r.certifications)

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                      topMargin=0.6 * inch, bottomMargin=0.6 * inch,
                      title=f"{r.name} - Resume", author=r.name).build(story)
    return buf.getvalue()
